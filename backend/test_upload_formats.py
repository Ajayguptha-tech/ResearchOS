"""Test PDF and DOCX upload and extraction."""
import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8', errors='replace')
sys.path.insert(0, ".")

from fastapi.testclient import TestClient
from app.main import app
from app.db.base import Base
from app.db.session import engine

Base.metadata.create_all(bind=engine)
client = TestClient(app)
tests = []

def test(name, fn):
    try:
        fn()
        tests.append((name, "PASS", ""))
    except AssertionError as e:
        tests.append((name, "FAIL", str(e)[:200]))
    except Exception as e:
        tests.append((name, "ERROR", str(e)[:200]))

# Login first
r = client.post("/api/v1/auth/login", json={"email": "test@example.com", "password": "TestPass123"})
assert r.status_code == 200, "Login failed"
token = r.json()["access_token"]
headers = {"Authorization": f"Bearer {token}"}

# Get project
r = client.get("/api/v1/projects/", headers=headers)
projects = r.json()
project_id = projects[0]["id"]

# TEST: Upload a real PDF
def test_pdf_upload():
    from fpdf import FPDF
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(200, 10, txt="Deep Learning for Network Security", new_x="LMARGIN", new_y="NEXT")
    pdf.cell(200, 10, txt="Abstract: This paper proposes a CNN-based intrusion detection system.", new_x="LMARGIN", new_y="NEXT")
    pdf.cell(200, 10, txt="Methodology: We use convolutional neural networks.", new_x="LMARGIN", new_y="NEXT")
    pdf.cell(200, 10, txt="Results: Achieved 95.3% accuracy on NSL-KDD dataset.", new_x="LMARGIN", new_y="NEXT")
    pdf_bytes = pdf.output()
    
    r = client.post(
        f"/api/v1/papers/upload?project_id={project_id}",
        files={"file": ("paper.pdf", bytes(pdf_bytes), "application/pdf")},
        headers=headers,
    )
    assert r.status_code == 201, f"PDF upload failed: {r.status_code} {r.text}"
    data = r.json()
    assert data["extracted_characters"] > 0, "No text extracted from PDF"
    assert "Deep Learning" in data.get("filename", "") or data["extracted_characters"] > 50
test("Upload PDF document", test_pdf_upload)

# TEST: Upload DOCX
def test_docx_upload():
    from docx import Document
    from io import BytesIO
    doc = Document()
    doc.add_heading("Machine Learning for Healthcare", level=1)
    doc.add_paragraph("This study examines the use of deep learning models for medical diagnosis.")
    doc.add_paragraph("We trained a ResNet-50 model on chest X-ray images.")
    doc.add_paragraph("Results show 92% accuracy in detecting pneumonia.")
    buf = BytesIO()
    doc.save(buf)
    
    r = client.post(
        f"/api/v1/papers/upload?project_id={project_id}",
        files={"file": ("healthcare.docx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        headers=headers,
    )
    assert r.status_code == 201, f"DOCX upload failed: {r.status_code} {r.text}"
    data = r.json()
    assert data["extracted_characters"] > 0, "No text extracted from DOCX"
    assert "Machine Learning" in data.get("filename", "") or data["extracted_characters"] > 50
test("Upload DOCX document", test_docx_upload)

# TEST: Upload CSV
def test_csv_upload():
    csv_data = "Name,Score,Method\nAlice,95,CNN\nBob,87,SVM\nCharlie,91,RF\n"
    r = client.post(
        f"/api/v1/papers/upload?project_id={project_id}",
        files={"file": ("results.csv", csv_data.encode(), "text/csv")},
        headers=headers,
    )
    assert r.status_code == 201, f"CSV upload failed: {r.status_code} {r.text}"
    data = r.json()
    assert data["extracted_characters"] > 0
test("Upload CSV document", test_csv_upload)

# TEST: List all documents in project
def test_list_docs():
    r = client.get(f"/api/v1/papers/documents/project/{project_id}", headers=headers)
    assert r.status_code == 200
    docs = r.json()
    assert len(docs) >= 3, f"Expected >= 3 docs, got {len(docs)}"
    filenames = [d["filename"] for d in docs]
    assert "paper.pdf" in filenames or "paper.pdf" in [d["filename"].lower() for d in docs]
    assert "healthcare.docx" in filenames or "healthcare.docx" in [d["filename"].lower() for d in docs]
test("List project documents after multiple uploads", test_list_docs)

# TEST: Run analysis with multiple documents
def test_multi_doc_analysis():
    r = client.post(
        f"/api/v1/research/analyze-local?idea=deep+learning+for+classification+and+healthcare&project_id={project_id}",
        headers=headers,
    )
    assert r.status_code == 200, f"Analysis failed: {r.status_code} {r.text}"
    result = r.json()
    assert result["analysis"]["papers_processed"] >= 2, f"Expected >= 2 papers, got {result['analysis']['papers_processed']}"
    # Check that documents are referenced
    docs_used = result.get("_documents_used", [])
    assert len(docs_used) >= 2, f"Expected >= 2 docs used, got {len(docs_used)}"
test("Analyze multiple uploaded documents", test_multi_doc_analysis)

# Print results
print()
print("=" * 60)
print("RESEARCHOS DOCUMENT UPLOAD & EXTRACTION TEST")
print("=" * 60)
passed = 0
failed = 0
for name, status, detail in tests:
    marker = "[PASS]" if status == "PASS" else "[FAIL]"
    suffix = f" -- {detail}" if detail else ""
    print(f"  {marker} {name}{suffix}")
    if status == "PASS":
        passed += 1
    else:
        failed += 1

print()
print(f"Results: {passed} passed, {failed} failed, {len(tests)} total")
if failed:
    print("OVERALL: FAIL")
    sys.exit(1)
else:
    print("OVERALL: ALL TESTS PASSED")
